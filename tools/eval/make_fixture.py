#!/usr/bin/env python
"""Generate the committed retrieval fixture (a 5-slide PPTX with fictional content).

Why a PPTX: the repository must ship a corpus that a *fresh clone* can index without any
private file, and every format the project reads needs a CJK font to be generated -
except PPTX, whose text lives in XML (no font embedding, no licence question) and whose
loader already assigns ``page_label`` = slide number, so the page-level metrics of
``tools/eval_retrieval.py`` stay meaningful.

The content is entirely fictional and mirrors the structure of the real corpus the
numbers were measured on: a cover, a scoring scheme, rules with numbers, a table, and a
penalty section. Regenerating it must be byte-stable-ish: the script always writes the
same text, so a new run changes nothing but the ZIP timestamps.

    python tools/eval/make_fixture.py            # rewrite the committed fixture
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "points-fixture.pptx"

#: (title, bullets) per slide. Keep the numbers here in sync with
#: ``tools/eval/fixture-points.json`` - the tests assert the terms really appear.
SLIDES: list[tuple[str, list[str]]] = [
    (
        "某学院学生综合素质积分管理办法（示例）",
        [
            "第一条 本办法用于示例评测，数据均为虚构。",
            "第二条 综合素质积分满分为 100 分。",
            "第三条 积分由五个部分组成：思想政治、学业成绩、实践服务、日常行为、群众评议。",
        ],
    ),
    (
        "学业成绩加分",
        [
            "第四条 英语四级通过加 4 分。",
            "第五条 英语六级通过加 6 分。",
            "第六条 计算机二级通过加 3 分。",
        ],
    ),
    (
        "实践服务加分",
        [
            "第七条 校内志愿服务每次加 2 分，累计上限 20 分。",
            "第八条 寒暑假社会实践每次加 3 分，累计上限 15 分。",
        ],
    ),
    (
        "理论考试与等级评定",
        [
            "第九条 理论考试 90 分以上加 2 分，缺考扣 3 分。",
            "第十条 等级对应分值见下表。",
        ],
    ),
    (
        "减分与说明",
        [
            "第十一条 宿舍卫生不合格每次扣 2 分。",
            "第十二条 加分材料弄虚作假的，取消当年度评优资格。",
            "第十三条 本办法由学生工作办公室负责解释。",
        ],
    ),
]

#: The table on slide 4, rendered by the loader as CSV.
TABLE = [
    ["等级", "分值"],
    ["优秀", "10 分"],
    ["良好", "8 分"],
    ["合格", "6 分"],
]


def build(path: Path = FIXTURE) -> Path:
    """Write the fixture and return its path (idempotent)."""
    presentation = Presentation()
    for index, (title, bullets) in enumerate(SLIDES, start=1):
        # layout 1 = "Title and Content": placeholders[0] is the title, [1] the body
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"{index}. {title}"
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for line in bullets[1:]:
            body.add_paragraph().text = line

        if index == 4:
            rows, cols = len(TABLE), len(TABLE[0])
            shape = slide.shapes.add_table(
                rows, cols, Inches(0.6), Inches(4.2), Inches(6.0), Inches(1.8)
            )
            for r in range(rows):
                for c in range(cols):
                    shape.table.cell(r, c).text = TABLE[r][c]

    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)
    return path


if __name__ == "__main__":
    written = build()
    print(f"wrote {written} ({written.stat().st_size} bytes, {len(SLIDES)} slides)")
