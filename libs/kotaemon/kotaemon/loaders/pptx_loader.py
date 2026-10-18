"""Local PowerPoint (.pptx) reader.

``.pptx`` used to go through ``unstructured``, which needs system ``libmagic`` and
fetches its own resources on first use — on this host that fails with a 403 and the
upload is silently unindexed (the same failure as ``.docx`` before MODIFICATIONS §2.8
and ``.xls`` before §2.9). ``python-pptx`` reads the file locally: no system library,
no network.

Reader behavior, mirroring :class:`DocxReader`:

* one Document per slide — text frames in shape order, tables rendered as CSV;
* speaker notes appended (they often carry the substance of a slide);
* ``page_label`` is the slide number, so the citation panel shows "[第 N 页]"
  like it does for PDFs, and the existing metadata contract is unchanged.
"""

from pathlib import Path
from typing import Any, List, Optional

import pandas as pd
from llama_index.core.readers.base import BaseReader

from kotaemon.base import Document


class PptxReader(BaseReader):
    """Read .pptx files with python-pptx (no system dependencies)."""

    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        try:
            import pptx  # noqa: F401
        except ImportError:
            raise ImportError(
                "python-pptx is not installed. "
                "Please install it using `pip install python-pptx`"
            )

    @staticmethod
    def _iter_shapes(shapes: Any):
        """Walk shapes, including the contents of grouped shapes."""
        for shape in shapes:
            # a group shape exposes nested `.shapes`; plain shapes do not
            if getattr(shape, "shapes", None) is not None:
                yield from PptxReader._iter_shapes(shape.shapes)
                continue
            yield shape

    def _table_to_csv(self, table: Any) -> str:
        """Render a pptx table as CSV, keeping the first row as the header."""
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        if not rows:
            return ""
        width = max(len(row) for row in rows)
        padded = [row + [""] * (width - len(row)) for row in rows]
        return pd.DataFrame(padded[1:], columns=padded[0]).to_csv(index=False).strip()

    def _slide_text(self, slide: Any) -> str:
        """All readable text of one slide: shapes, tables and notes."""
        parts: List[str] = []
        for shape in self._iter_shapes(slide.shapes):
            if getattr(shape, "has_table", False) and shape.has_table:
                table_text = self._table_to_csv(shape.table)
                if table_text:
                    parts.append(table_text)
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)

        if getattr(slide, "has_notes_slide", False):
            notes_frame = slide.notes_slide.notes_text_frame
            notes_text = notes_frame.text.strip() if notes_frame else ""
            if notes_text:
                parts.append(f"[notes] {notes_text}")

        return "\n".join(part for part in parts if part).strip()

    def load_data(
        self, file_path: Path, extra_info: Optional[dict] = None, **kwargs: Any
    ) -> List[Document]:
        """Load data using python-pptx.

        Args:
            file_path: path to the .pptx file
            extra_info: metadata merged into every returned Document

        Returns:
            one Document per slide that carries text
        """
        from pptx import Presentation

        file_path = Path(file_path).resolve()
        extra_info = extra_info or {}

        presentation = Presentation(str(file_path))
        documents: List[Document] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            text = self._slide_text(slide)
            if not text:
                continue  # an image-only slide contributes nothing to retrieval
            documents.append(
                Document(
                    text=text,
                    metadata={
                        "page_label": slide_number,
                        "slide_number": slide_number,
                        **extra_info,
                    },
                )
            )

        if not documents:
            # keep the file listed even when it holds no text, so the file list and
            # the doctor check can tell "empty deck" from "failed parse"
            documents.append(
                Document(
                    text="",
                    metadata={"page_label": 1, "slide_number": 1, **extra_info},
                )
            )
        return documents
